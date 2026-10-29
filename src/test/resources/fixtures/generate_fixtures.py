# -*- coding: utf-8 -*-
"""生成 IMM 按页切图集成测试用的 10 页 docx 样例。

用法: python generate_fixtures.py
产出: imm-sample-10p.docx（同目录）
说明: .doc 版本由 Word COM 另存生成，见同目录 generate.md。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "imm-sample-10p.docx"

SECTIONS = [
    ("第一章 项目概述", "本项目为宿舍楼消防设备采购项目，涵盖独立式烟感报警器的供货、安装与调试。"),
    ("第二章 招标范围", "供应商需提供符合国标要求的烟感报警器一千套，并负责现场安装指导。"),
    ("第三章 投标人资格要求", "投标人须具有独立法人资格，具备同类项目业绩不少于三项。"),
    ("第四章 技术要求", "报警器需支持光电感应原理，工作温度范围零下十度至五十度。"),
    ("第五章 商务要求", "交货期为合同签订后三十日内，质保期不少于两年。"),
    ("第六章 评标办法", "采用综合评分法，技术分占百分之六十，商务分占百分之四十。"),
    ("第七章 投标文件格式", "投标文件应包含商务标、技术标与报价表三部分，均需加盖公章。"),
    ("第八章 开标与评标程序", "开标时间为递交截止时间后一个工作日内，评标委员会由五人组成。"),
    ("第九章 合同主要条款", "合同价款采用固定总价方式，支付方式为到货验收后支付百分之九十。"),
    ("第十章 附则", "本招标文件的解释权归招标人所有，未尽事宜在答疑纪要中补充。"),
]


def build_page(doc, title, paragraph, last=False):
    heading = doc.add_heading(title, level=1)
    for run in heading.runs:
        run.font.name = "宋体"
    doc.add_paragraph(paragraph)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "条目"
    table.cell(0, 1).text = "内容"
    table.cell(0, 2).text = "备注"
    table.cell(1, 0).text = title
    table.cell(1, 1).text = paragraph[:20]
    table.cell(2, 0).text = "页示例"
    table.cell(2, 1).text = "本页为按页切图测试样例页"
    table.cell(2, 2).text = "IMM OfficeConversion"
    for row in table.rows:
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10.5)
    if not last:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def main():
    doc = Document()
    for index, (title, paragraph) in enumerate(SECTIONS):
        build_page(doc, title, paragraph, last=index == len(SECTIONS) - 1)
    doc.save(OUTPUT)
    print(f"已生成: {OUTPUT} ({len(SECTIONS)} 页)")


if __name__ == "__main__":
    main()
